from io import BytesIO
from datetime import datetime

import pandas as pd
from flask import Blueprint, jsonify, request, send_file
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

from models.student_model import get_filtered_student_data, get_grade_class_summary
from routes.auth_routes import login_required


export_bp = Blueprint('export', __name__)


@export_bp.post('/api/export-summary')
@login_required
def export_summary():
    age = request.form.get('age', '')
    age_1 = request.form.get('age_1', '')
    age_2 = request.form.get('age_2', '')
    class_name = request.form.get('class', 'all')
    gender = request.form.get('gender', 'all')
    grade = request.form.get('grade', 'all')
    date = request.form.get('date') or datetime.now().strftime('%Y-%m-%d')
    filename = request.form.get('file_name') or 'grade_class_summary'
    summary = get_grade_class_summary(
        date, class_name, gender, grade, age, age_1, age_2,
    )
    download_data = [
        {
            'Grade': row['grade'],
            'Class': row['class'],
            'Male': row['male'],
            'Female': row['female'],
            'Total': row['total'],
        }
        for row in summary['rows']
    ]
    download_data.append({
        'Grade': 'Grand Total',
        'Class': '',
        'Male': summary['grand_total']['male'],
        'Female': summary['grand_total']['female'],
        'Total': summary['grand_total']['total'],
    })

    output = BytesIO()
    df = pd.DataFrame(download_data, columns=['Grade', 'Class', 'Male', 'Female', 'Total'])
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)

    output.seek(0)
    workbook = load_workbook(output)
    worksheet = workbook.active
    worksheet.merge_cells(start_row=worksheet.max_row, start_column=1, end_row=worksheet.max_row, end_column=2)

    widths = {'A': 18, 'B': 14, 'C': 12, 'D': 12, 'E': 12}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows():
        worksheet.row_dimensions[row[0].row].height = 20
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center')

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    grand_total_row = worksheet.max_row
    for cell in worksheet[grand_total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D9EAD3', end_color='D9EAD3', fill_type='solid')

    final_output = BytesIO()
    workbook.save(final_output)
    final_output.seek(0)
    return send_file(
        final_output,
        as_attachment=True,
        download_name=f'{filename}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@export_bp.post('/api/export')
@login_required
def export_data():
    age = request.form.get('age', '')
    age_1 = request.form.get('age_1', '')
    age_2 = request.form.get('age_2', '')
    class_name = request.form.get('class', 'all')
    gender = request.form.get('gender', 'all')
    grade = request.form.get('grade', 'all')
    date = request.form.get('date') or datetime.now().strftime('%Y-%m-%d')
    filename = request.form.get('file_name') or 'student_export'
    file_type = request.form.get('file_type', 'normal')
    school_name = request.form.get('school_name') or ''

    data = get_filtered_student_data(date, age, age_1, age_2, class_name, gender, grade)
    output = BytesIO()

    if file_type == 'normal':
        download_data = [
            {
                'စဉ်': row[0],
                'ကျောင်းဝင်အမှတ်': row[1],
                'နာမည်': row[2],
                'အဖေနာမည်': row[3],
                'ကျားမ': row[4],
                'မွေးနေ့': row[5],
                'class': row[6],
                'Grade': row[7],
                'အသက်': row[8],
            }
            for row in data
        ]
        df = pd.DataFrame(download_data)

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False)

        output.seek(0)
        workbook = load_workbook(output)
        worksheet = workbook.active

        worksheet.column_dimensions['A'].width = 6
        worksheet.column_dimensions['B'].width = 18
        worksheet.column_dimensions['C'].width = 26
        worksheet.column_dimensions['D'].width = 26
        worksheet.column_dimensions['E'].width = 8
        worksheet.column_dimensions['F'].width = 12
        worksheet.column_dimensions['G'].width = 9
        worksheet.column_dimensions['H'].width = 9
        worksheet.column_dimensions['I'].width = 24

        for row in worksheet.iter_rows():
            worksheet.row_dimensions[row[0].row].height = 20

        for cell in worksheet[1]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for row in worksheet.iter_rows(min_row=2):
            row[0].alignment = Alignment(horizontal='center', vertical='center')
            row[1].alignment = Alignment(horizontal='center', vertical='center')
            row[2].alignment = Alignment(horizontal='left', vertical='center')
            row[3].alignment = Alignment(horizontal='left', vertical='center')
            row[4].alignment = Alignment(horizontal='center', vertical='center')
            row[5].alignment = Alignment(horizontal='center', vertical='center')
            row[6].alignment = Alignment(horizontal='center', vertical='center')
            row[7].alignment = Alignment(horizontal='center', vertical='center')
            row[8].alignment = Alignment(horizontal='center', vertical='center')

        final_output = BytesIO()
        workbook.save(final_output)
        final_output.seek(0)
        return send_file(
            final_output,
            as_attachment=True,
            download_name=f'{filename}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    if file_type == 'normalized':
        no, school_names, rolls, names = [], [], [], []
        fathers, genders, dobs, reasons = [], [], [], []
        for row in data:
            no.append(row[0])
            school_names.append(school_name)
            rolls.append(row[1])
            names.append(row[2])
            fathers.append(row[3])
            genders.append(row[4])
            dobs.append(row[5])
            reasons.append('')

        download_data = {
            'စဉ်': no,
            'ကျောင်းအမည်': school_names,
            'ကျောင်းဝင်အမှတ်': rolls,
            'အမည်': names,
            'အဖေအမည်': fathers,
            'ကျား/မ': genders,
            'မွေးသက္ကရာဇ်': dobs,
            'မှတ်ချက်': reasons,
        }
        df = pd.DataFrame(download_data)

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False)

        output.seek(0)
        workbook = load_workbook(output)
        worksheet = workbook.active

        thin = Side(style='thin')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = border

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        worksheet.insert_rows(1, amount=4)
        worksheet['H1'] = 'ပူးတွဲ(က)'
        worksheet['A2'] = 'အခြေခံပညာဦးစီးဌာန'
        worksheet['A3'] = 'ရန်ကုန်တိုင်းဒေသကြီးအင်းစိန်ခရိုင်လှိုင်သာယာ(အနောက်ပိုင်း)မြို့နယ်'
        worksheet['A4'] = '၂၀၂၆ - ၂၀၂၇ ပညာသင်နှစ်၊ အခြေခံပညာ မူလတန်းအဆင့် Grade5 စာမေးပွဲဖြေဆိုသူစာရင်းပေါင်းချုပ်'

        worksheet.merge_cells('A2:H2')
        worksheet.merge_cells('A3:H3')
        worksheet.merge_cells('A4:H4')

        for row in range(1, 5):
            cell = worksheet[f'A{row}']
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = Font(bold=True, size=13)

        worksheet['H1'].font = Font(bold=True, size=13)
        worksheet['H1'].alignment = Alignment(horizontal='center', vertical='center')
        header_fill = PatternFill(start_color='D9EAD3', end_color='D9EAD3', fill_type='solid')

        for cell in worksheet[5]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = header_fill

        worksheet.column_dimensions['A'].width = 6
        worksheet.column_dimensions['B'].width = 14
        worksheet.column_dimensions['C'].width = 18
        worksheet.column_dimensions['D'].width = 26
        worksheet.column_dimensions['E'].width = 26
        worksheet.column_dimensions['F'].width = 8
        worksheet.column_dimensions['G'].width = 15
        worksheet.column_dimensions['H'].width = 24

        for row in worksheet.iter_rows():
            worksheet.row_dimensions[row[0].row].height = 23
        for row in worksheet.iter_rows(min_row=5):
            worksheet.row_dimensions[row[0].row].height = 18.5
        for row in worksheet.iter_rows():
            for cell in row:
                cell.font = Font(name='Pyidaungsu', size=13, bold=True)
        for row in worksheet.iter_rows(min_row=5):
            for cell in row:
                cell.font = Font(name='Pyidaungsu', size=11)

        alignments = ['center', 'left', 'center', 'left', 'left', 'center', 'center', 'left']
        for row in worksheet.iter_rows(min_row=6):
            for cell, alignment in zip(row, alignments):
                cell.alignment = Alignment(horizontal=alignment, vertical='center')

        for row in worksheet.iter_rows(min_row=5):
            for cell in row:
                cell.border = border

        worksheet.freeze_panes = 'A6'
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.orientation = 'portrait'
        worksheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5)
        worksheet.print_title_rows = '1:5'
        worksheet.print_options.horizontalCentered = True

        data_start_row = 5
        rows_per_page = 22
        current_row = data_start_row + rows_per_page
        while current_row <= worksheet.max_row:
            worksheet.row_breaks.append(Break(id=current_row))
            current_row += rows_per_page

        final_output = BytesIO()
        workbook.save(final_output)
        final_output.seek(0)
        return send_file(
            final_output,
            as_attachment=True,
            download_name=f'{filename}_normalized.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    return jsonify({'error': 'Unsupported file type.'}), 400